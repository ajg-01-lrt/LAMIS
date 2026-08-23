"""Regression: File Processing → BoM build must heal a Summary that's
missing rows for device tabs already in the workbook.

Bug from the 2026-06-08 troubleshooting workbook the operator sent:
the inventory file had 11 device tabs but Summary only listed one of
them (separate inventory-builder bug fixed in build_psi_report_workbook).
File Processing → BoM then ran ``_extract_summary``, got 1 site, and
produced a BoM with a single site column — the operator's 10 other
shelves silently dropped out of the aggregate.

The fix scans every non-Summary, non-BOM sheet for device-tab shape
(yields BOM entries via the standard parser), and for any tab not
covered by Summary, adds it to ``summary_items`` and rewrites the
Summary sheet so the workbook itself gets healed.
"""
from __future__ import annotations
import os
import shutil
import tempfile
import unittest
from unittest.mock import MagicMock

import openpyxl

from gui.bom_frame import BomFrame
from gui.workbook_builder import WorkbookBuilder


def _make_builder() -> WorkbookBuilder:
    db = MagicMock()
    db.db_path = ":memory:"
    db.lookup_part.return_value = ""
    return WorkbookBuilder(
        db_cache=db, template_path="", packing_slip_template="",
    )


def _add_device_tab(
    wb,
    tab_name: str,
    ip: str,
    system_name: str,
    parts: list,
) -> None:
    """Add a device tab in the standard layout
    (B=Name, C=Type, D=Part, E=Serial, F=Description from row 15)
    with F5=IP, F6=System Name metadata."""
    ws = wb.create_sheet(tab_name)
    ws["F5"] = ip
    ws["F6"] = system_name
    ws["F7"] = "Test System"
    for idx, (name, ptype, part, serial, desc) in enumerate(parts):
        r = 15 + idx
        ws[f"B{r}"] = name
        ws[f"C{r}"] = ptype
        ws[f"D{r}"] = part
        ws[f"E{r}"] = serial
        ws[f"F{r}"] = desc


def _add_summary_row(
    summary_ws, row: int, idx: int, ip: str, device_name: str, link_tab: str,
) -> None:
    summary_ws[f"B{row}"] = idx
    summary_ws[f"C{row}"] = ip
    summary_ws[f"D{row}"] = device_name
    summary_ws[f"D{row}"].hyperlink = f"#'{link_tab}'!A1"


class TestBomFrameHealsTruncatedSummary(unittest.TestCase):

    def setUp(self):
        self.tmp_dir = tempfile.mkdtemp(prefix="bom_heal_")

    def tearDown(self):
        shutil.rmtree(self.tmp_dir, ignore_errors=True)

    def _build_truncated_workbook(self) -> str:
        """Build a workbook with 3 device tabs but Summary listing only
        the first — exactly the shape ``build_psi_report_workbook``
        produced before the matching builder-side fix landed."""
        wb = openpyxl.Workbook()
        summary = wb.active
        summary.title = "Summary"

        # Standard Summary header (rows 5-9 carry labels in the real
        # template; only the parts the BoM build reads matter here).
        summary["B5"] = "Customer"
        summary["D5"] = "Project"
        summary["B7"] = "Test Cust"
        summary["D7"] = "Test Proj"
        summary["F5"] = "Device Count"
        summary["F7"] = 1
        summary["B9"] = "#"
        summary["C9"] = "IP Address"
        summary["D9"] = "Device Name"
        summary["E9"] = "Asset Tag"

        parts = [
            ("Shelf 1", "Shelf", "NTK-CHASSIS-1", "M9A03562",
             "6500-R4 Chassis"),
            ("Card 1", "Card", "NTK-CARD-1", "C-001", "Some Card"),
        ]
        _add_device_tab(
            wb, "device_one_bb_net_apple_com", "10.0.0.1",
            "device-one.bb.net.apple.com", parts,
        )
        _add_device_tab(
            wb, "device_two_bb_net_apple_com", "10.0.0.5",
            "device-two.bb.net.apple.com", parts,
        )
        _add_device_tab(
            wb, "device_three_bb_net_apple_com", "10.0.0.6",
            "device-three.bb.net.apple.com", parts,
        )

        # Summary only lists device ONE — devices two and three are
        # orphans that the BoM build must discover and add.
        _add_summary_row(
            summary, row=10, idx=1,
            ip="10.0.0.1",
            device_name="device-one.bb.net.apple.com",
            link_tab="device_one_bb_net_apple_com",
        )

        path = os.path.join(self.tmp_dir, "truncated.xlsx")
        wb.save(path)
        return path

    def test_augment_adds_orphan_device_tabs_to_summary_items(self):
        path = self._build_truncated_workbook()
        wb = openpyxl.load_workbook(path)
        builder = _make_builder()

        # Bypass __init__ — the BomFrame widget needs a parent + gui,
        # but the heal helper is self-contained. NEW the instance and
        # stub the log sink.
        frame = BomFrame.__new__(BomFrame)
        frame._append_log = lambda _msg: None  # type: ignore[assignment]

        sheetname_set = set(wb.sheetnames)
        summary_items, display_to_tab = BomFrame._extract_summary(
            wb["Summary"], sheetname_set,
        )
        # Truncated Summary: only one site initially.
        self.assertEqual(len(summary_items), 1)

        added = frame._augment_summary_from_device_tabs(
            wb, builder, summary_items, display_to_tab,
        )
        # Two orphan tabs picked up.
        self.assertEqual(added, 2)
        # summary_items mutated in place to include all three.
        self.assertEqual(len(summary_items), 3)
        names = {item[1] for item in summary_items}
        self.assertIn("device-one.bb.net.apple.com", names)
        self.assertIn("device-two.bb.net.apple.com", names)
        self.assertIn("device-three.bb.net.apple.com", names)

    def test_summary_sheet_rewritten_with_full_site_list(self):
        path = self._build_truncated_workbook()
        wb = openpyxl.load_workbook(path)
        builder = _make_builder()

        frame = BomFrame.__new__(BomFrame)
        frame._append_log = lambda _msg: None

        sheetname_set = set(wb.sheetnames)
        summary_items, display_to_tab = BomFrame._extract_summary(
            wb["Summary"], sheetname_set,
        )
        frame._augment_summary_from_device_tabs(
            wb, builder, summary_items, display_to_tab,
        )

        # Workbook's Summary sheet now lists all three devices.
        summary = wb["Summary"]
        d_col = [summary.cell(r, 4).value for r in range(10, 14)]
        non_empty = [v for v in d_col if v]
        self.assertEqual(
            len(non_empty), 3,
            f"Healed Summary should list all 3 devices; got {d_col!r}",
        )
        self.assertEqual(summary["F7"].value, 3)

    def test_already_complete_summary_is_left_alone(self):
        """Sanity: when Summary already lists every device tab, the
        heal step must NOT add duplicates or churn the sheet."""
        # Build a workbook where Summary already covers both tabs.
        wb = openpyxl.Workbook()
        summary = wb.active
        summary.title = "Summary"
        summary["B9"], summary["C9"] = "#", "IP Address"
        summary["D9"] = "Device Name"
        parts = [("Shelf", "Shelf", "P-1", "S-1", "Desc")]
        _add_device_tab(wb, "dev_a", "10.0.0.1", "dev-a", parts)
        _add_device_tab(wb, "dev_b", "10.0.0.2", "dev-b", parts)
        _add_summary_row(summary, 10, 1, "10.0.0.1", "dev-a", "dev_a")
        _add_summary_row(summary, 11, 2, "10.0.0.2", "dev-b", "dev_b")
        path = os.path.join(self.tmp_dir, "complete.xlsx")
        wb.save(path)

        wb = openpyxl.load_workbook(path)
        builder = _make_builder()
        frame = BomFrame.__new__(BomFrame)
        frame._append_log = lambda _msg: None

        summary_items, display_to_tab = BomFrame._extract_summary(
            wb["Summary"], set(wb.sheetnames),
        )
        before = len(summary_items)
        added = frame._augment_summary_from_device_tabs(
            wb, builder, summary_items, display_to_tab,
        )
        self.assertEqual(added, 0, "Complete Summary must add nothing")
        self.assertEqual(len(summary_items), before)


if __name__ == "__main__":
    unittest.main()
