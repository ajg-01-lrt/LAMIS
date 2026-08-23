"""Tests for the live Asset Tag XLOOKUP formula on device tabs.

Operator's request: rather than copying Asset Tag values from Summary
to device tabs as literals (the existing ``propagate_asset_tags_to_tabs``
behavior, which required a BoM rebuild to refresh), every device
tab's G15 should now carry a live XLOOKUP that pulls from Summary E
keyed on the device hostname. Edit Summary once, every device tab
follows automatically.
"""
from __future__ import annotations
import os
import shutil
import tempfile
import unittest
from unittest.mock import MagicMock

import openpyxl
import pandas as pd

from gui.workbook_builder import WorkbookBuilder


def _make_builder() -> WorkbookBuilder:
    db = MagicMock()
    db.db_path = ":memory:"
    db.lookup_part.return_value = ""
    return WorkbookBuilder(
        db_cache=db, template_path="", packing_slip_template="",
    )


EXPECTED_FORMULA = '=XLOOKUP(F6,Summary!$D:$D,Summary!$E:$E,"",0,1)'


class TestWriteAssetTagFormulaHelper(unittest.TestCase):
    """Unit tests for the new ``write_asset_tag_formula_to_chassis_row``
    helper — the workhorse called from every inventory builder."""

    def setUp(self):
        self.builder = _make_builder()
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active

    def test_writes_formula_to_g15_by_default(self):
        wrote = self.builder.write_asset_tag_formula_to_chassis_row(self.ws)
        self.assertTrue(wrote)
        self.assertEqual(self.ws["G15"].value, EXPECTED_FORMULA)

    def test_formula_keys_on_hostname_in_f6(self):
        # F6 carries the device's system_name in every ATLAS-built
        # device tab — that's the lookup key the formula uses against
        # Summary's Device Name column.
        self.builder.write_asset_tag_formula_to_chassis_row(self.ws)
        formula = self.ws["G15"].value
        self.assertIn("F6", formula)
        self.assertIn("Summary!$D:$D", formula)
        self.assertIn("Summary!$E:$E", formula)

    def test_existing_formula_preserved_by_default(self):
        # Idempotent: a prior build already planted the formula —
        # don't re-write it (avoids dirtying the workbook for no
        # reason on append-mode rebuilds).
        self.ws["G15"] = '=XLOOKUP(F6,Summary!$D:$D,Summary!$E:$E,"",0,1)'
        wrote = self.builder.write_asset_tag_formula_to_chassis_row(self.ws)
        self.assertFalse(wrote)

    def test_literal_value_overwritten_by_formula(self):
        # An older workbook may have a literal asset tag at G15
        # (typed directly or stamped by the prior propagate flow).
        # New builds replace it with the formula so the live link
        # works going forward.
        self.ws["G15"] = "OLD_LITERAL_TAG_123"
        wrote = self.builder.write_asset_tag_formula_to_chassis_row(self.ws)
        self.assertTrue(wrote)
        self.assertEqual(self.ws["G15"].value, EXPECTED_FORMULA)

    def test_force_overwrites_existing_formula(self):
        self.ws["G15"] = '=SOME_OTHER_FORMULA()'
        wrote = self.builder.write_asset_tag_formula_to_chassis_row(
            self.ws, force=True,
        )
        self.assertTrue(wrote)
        self.assertEqual(self.ws["G15"].value, EXPECTED_FORMULA)


class TestPropagateLeavesFormulaCellAlone(unittest.TestCase):
    """``write_asset_tag_to_chassis_row`` is the legacy propagate
    primitive — it MUST skip cells that hold a formula now, or the
    propagate pass at BoM-build time would silently replace the live
    XLOOKUP with a frozen literal."""

    def setUp(self):
        self.builder = _make_builder()
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        # Plant a chassis-shaped row at row 15 so the chassis-row
        # finder targets it.
        self.ws["B15"] = "Shelf 1"
        self.ws["C15"] = "Shelf"

    def test_skips_when_g15_is_formula(self):
        self.ws["G15"] = '=XLOOKUP(F6,Summary!$D:$D,Summary!$E:$E,"",0,1)'
        row = self.builder.write_asset_tag_to_chassis_row(self.ws, "NEW_TAG")
        self.assertEqual(row, 15)
        # Formula preserved.
        self.assertTrue(self.ws["G15"].value.startswith("="))

    def test_writes_literal_when_g15_was_empty(self):
        row = self.builder.write_asset_tag_to_chassis_row(self.ws, "TAG_42")
        self.assertEqual(row, 15)
        self.assertEqual(self.ws["G15"].value, "TAG_42")

    def test_overwrites_existing_literal(self):
        # Sanity: the formula-skip guard doesn't accidentally protect
        # legitimate operator-typed literals from being refreshed by
        # propagate.
        self.ws["G15"] = "STALE_TAG"
        self.builder.write_asset_tag_to_chassis_row(self.ws, "FRESH_TAG")
        self.assertEqual(self.ws["G15"].value, "FRESH_TAG")


class TestBuildReportWorkbookPlantsFormula(unittest.TestCase):
    """End-to-end: a fresh inventory report from the standard builder
    has every device tab's G15 holding the live XLOOKUP."""

    def setUp(self):
        self.tmp_dir = tempfile.mkdtemp(prefix="asset_tag_formula_")
        template_path = os.path.join(self.tmp_dir, "report_template.xlsx")
        tpl = openpyxl.Workbook()
        tpl.active.title = "Customer-project"
        tpl.create_sheet("Summary")
        # Device template carries G14="Asset Tag" header already in
        # the real install — replicate so the test sheet looks right.
        dev_tpl = tpl.create_sheet("DeviceTemplate")
        dev_tpl["G14"] = "Asset Tag"
        tpl.save(template_path)

        db = MagicMock()
        db.db_path = ":memory:"
        db.lookup_part.return_value = ""
        self.builder = WorkbookBuilder(
            db_cache=db, template_path=template_path,
            packing_slip_template="",
        )
        self.template_path = template_path

    def tearDown(self):
        shutil.rmtree(self.tmp_dir, ignore_errors=True)

    def _device_data(self, system_name: str) -> dict:
        df = pd.DataFrame([{
            "System Name": system_name,
            "System Type": "Test System",
            "Type": "Shelf",
            "Information Type": "Shelf",
            "Part Number": "P-1",
            "Serial Number": "S-1",
            "Description": "Test chassis",
            "Name": "Main Shelf",
            "Source": "LAN",
        }])
        return {"main": df}

    def test_g15_holds_formula_on_each_device_tab(self):
        out_path = os.path.join(self.tmp_dir, "out.xlsx")
        self.builder.build_report_workbook(
            {
                "10.0.0.1": self._device_data("device-one.example.com"),
                "10.0.0.2": self._device_data("device-two.example.com"),
            },
            out_path,
            customer="Test", project="Proj",
            customer_po="PO", sales_order="SO",
            append_mode=False,
        )
        # Reload preserving formulas.
        wb = openpyxl.load_workbook(out_path)
        device_tabs = [
            n for n in wb.sheetnames
            if n not in ("Summary", "Inventory by Site")
            and not n.startswith("DeviceTemplate")
            and not n.startswith("Customer")
        ]
        self.assertEqual(
            len(device_tabs), 2,
            f"Expected 2 device tabs, got {device_tabs!r}",
        )
        for tab in device_tabs:
            g15 = wb[tab]["G15"].value
            self.assertTrue(
                isinstance(g15, str) and g15.startswith("="),
                f"Tab {tab!r} G15 should be a formula, got {g15!r}",
            )
            # Hostname-keyed and pointing at Summary E.
            self.assertIn("F6", g15)
            self.assertIn("Summary", g15)


if __name__ == "__main__":
    unittest.main()
