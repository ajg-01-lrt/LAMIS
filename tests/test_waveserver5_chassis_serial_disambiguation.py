"""Tests for the ``WS5_1`` → ``WS5-<serial>`` disambiguation.

Operators upgrade Waveserver 5s through ATLAS's flow, which hard-codes
the post-provision hostname to ``WS5_1``. Inventorying multiple WS5s
in one workbook then collapses every chassis onto the same tab, BoM
column, and Summary row — because they all reported the same hostname.

The fix piggybacks on the existing ``_FACTORY_DEFAULT_HOSTNAMES`` map:
when a device reports a listed hostname, ATLAS rewrites the system
name to ``<prefix>-<chassis-serial>`` BEFORE sheet creation, so two
WS5s with the same TID become e.g. ``WS5-M9A03562`` and
``WS5-M9A12345`` — each with its own tab.
"""
from __future__ import annotations
import os
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock

import openpyxl
import pandas as pd

from gui.workbook_builder import (
    WorkbookBuilder,
    _FACTORY_DEFAULT_HOSTNAMES,
    _chassis_serial_from_df,
)


class TestFactoryDefaultMapIncludesWaveserver5(unittest.TestCase):

    def test_ws5_1_listed(self):
        # ATLAS's upgrade flow hard-codes the post-provision hostname,
        # so the map MUST recognize it or every WS5 collapses onto one
        # sheet in multi-shelf workbooks.
        self.assertEqual(_FACTORY_DEFAULT_HOSTNAMES.get("ws5_1"), "WS5")

    def test_waveserver_5_listed(self):
        # Un-provisioned devices report ``Waveserver-5`` literally.
        self.assertEqual(_FACTORY_DEFAULT_HOSTNAMES.get("waveserver-5"), "WS5")

    def test_lookup_is_case_insensitive_via_lower_keys(self):
        # The map is meant to be looked up after ``.strip().lower()``
        # on the hostname (see build_report_workbook). Keys must
        # already be lowercase or the lookup misses on real input.
        for key in _FACTORY_DEFAULT_HOSTNAMES:
            self.assertEqual(
                key, key.lower(),
                f"_FACTORY_DEFAULT_HOSTNAMES key {key!r} must be lowercase "
                f"(callers always lowercase the hostname before lookup)",
            )

    def test_existing_rls_entry_preserved(self):
        # Regression guard — the RLS rewrite is load-bearing and must
        # not be lost while extending the map.
        self.assertEqual(_FACTORY_DEFAULT_HOSTNAMES.get("rls"), "RLS")


class TestChassisSerialFromInformationType(unittest.TestCase):
    """``_chassis_serial_from_df`` originally only checked the ``Type``
    column. Waveserver 5 stores the descriptive model name there
    (``"Waveserver 5 Chassis"``) and tags the BoM classifier in the
    ``Information Type`` column (``"Shelf"``). The helper must accept
    either form."""

    def test_picks_serial_via_information_type_shelf(self):
        df = pd.DataFrame([
            {
                "Type": "Waveserver 5 Chassis",
                "Information Type": "Shelf",
                "Serial Number": "M9A03562",
            },
            {
                "Type": "Waveserver 5 Fan Module",
                "Information Type": "Component",
                "Serial Number": "FAN-SERIAL-1",
            },
        ])
        self.assertEqual(_chassis_serial_from_df(df), "M9A03562")

    def test_picks_serial_via_information_type_chassis(self):
        df = pd.DataFrame([
            {
                "Type": "Waveserver 5 Chassis",
                "Information Type": "Chassis",
                "Serial Number": "M9A12345",
            },
        ])
        self.assertEqual(_chassis_serial_from_df(df), "M9A12345")

    def test_old_shape_still_works(self):
        # Existing scripts (RLS, SAR, IXR) put the classifier in Type.
        # The helper must keep accepting that.
        df = pd.DataFrame([
            {"Type": "Shelf", "Serial Number": "OLD-SHAPE-SERIAL"},
        ])
        self.assertEqual(_chassis_serial_from_df(df), "OLD-SHAPE-SERIAL")

    def test_skips_chassis_row_with_empty_serial(self):
        # Falls through to subsequent chassis rows if the first is
        # missing its serial — important for partial reads.
        df = pd.DataFrame([
            {"Type": "Shelf", "Information Type": "Shelf", "Serial Number": ""},
            {"Type": "Card", "Information Type": "Card", "Serial Number": "CARD-1"},
            {"Type": "Shelf", "Information Type": "Shelf", "Serial Number": "ACTUAL-SERIAL"},
        ])
        self.assertEqual(_chassis_serial_from_df(df), "ACTUAL-SERIAL")

    def test_returns_empty_when_no_chassis(self):
        df = pd.DataFrame([
            {"Type": "Card", "Information Type": "Card", "Serial Number": "C-1"},
        ])
        self.assertEqual(_chassis_serial_from_df(df), "")


class TestTwoWaveserver5sGetDistinctTabs(unittest.TestCase):
    """End-to-end: two WS5 devices reporting the same ``WS5_1``
    hostname but different chassis serials must land in separate
    tabs / BoM columns / Summary rows."""

    def setUp(self):
        self.tmp_dir = tempfile.mkdtemp(prefix="ws5_disambig_")
        template_path = os.path.join(self.tmp_dir, "report_template.xlsx")
        tpl = openpyxl.Workbook()
        tpl.active.title = "Customer-project"
        tpl.create_sheet("Summary")
        tpl.create_sheet("DeviceTemplate")
        tpl.save(template_path)

        db = MagicMock()
        db.db_path = ":memory:"
        db.lookup_part.return_value = ""
        self.builder = WorkbookBuilder(
            db_cache=db,
            template_path=template_path,
            packing_slip_template="",
        )
        self.template_path = template_path

    def tearDown(self):
        shutil.rmtree(self.tmp_dir, ignore_errors=True)

    def _ws5_data(self, chassis_serial: str, slot1_serial: str) -> dict:
        """Build the per-IP data_dict shape that build_report_workbook
        expects — a dict of category → DataFrame. Mirrors what
        Ciena_Waveserver5.process_outputs would produce."""
        rows = [
            {
                "System Name": "WS5_1",
                "System Type": "Waveserver 5",
                "Type": "Waveserver 5 Chassis",
                "Information Type": "Shelf",
                "Part Number": "186-3001-900",
                "Serial Number": chassis_serial,
                "Description": "Waveserver 5 Chassis",
                "Name": "Chassis",
                "Source": "LAN",
            },
            {
                "System Name": "WS5_1",
                "System Type": "Waveserver 5",
                "Type": "Waveserver WL5e 2x800G ... Module",
                "Information Type": "Card",
                "Part Number": "186-3161-900",
                "Serial Number": slot1_serial,
                "Description": "WL5e module",
                "Name": "Slot 1",
                "Source": "LAN",
            },
        ]
        return {"main": pd.DataFrame(rows)}

    @staticmethod
    def _ws5_device_sheets(wb) -> list:
        # Sheet titles strip non-alphanumeric chars to ``_`` via
        # ``make_unique_sheet_title`` — the rewritten ``WS5-<serial>``
        # name lands as ``WS5_<serial>`` on the tab strip. The dash
        # form is preserved in F6 and in BoM column headers.
        return sorted(
            n for n in wb.sheetnames
            if n.lower().startswith("ws5_")
        )

    def test_two_ws5s_same_tid_get_separate_sheets(self):
        outputs = {
            "10.0.0.10": self._ws5_data("M9A03562", "M9A16485"),
            "10.0.0.11": self._ws5_data("M9A12345", "M9A2DCD6"),
        }
        out_path = os.path.join(self.tmp_dir, "out.xlsx")
        self.builder.build_report_workbook(
            outputs, out_path,
            customer="Apple", project="QAS-UFG",
            customer_po="PO", sales_order="SO",
            append_mode=False,
        )
        wb = openpyxl.load_workbook(out_path)
        device_tabs = self._ws5_device_sheets(wb)
        self.assertEqual(
            len(device_tabs), 2,
            f"Two WS5s with same TID should get two distinct tabs, "
            f"got {device_tabs!r} (all sheets: {wb.sheetnames})",
        )
        # Each tab name contains its chassis serial — that's how the
        # operator tells them apart at a glance.
        joined = " ".join(device_tabs)
        self.assertIn("M9A03562", joined)
        self.assertIn("M9A12345", joined)

    def test_summary_lists_both_devices(self):
        outputs = {
            "10.0.0.10": self._ws5_data("M9A03562", "M9A16485"),
            "10.0.0.11": self._ws5_data("M9A12345", "M9A2DCD6"),
        }
        out_path = os.path.join(self.tmp_dir, "summary_check.xlsx")
        self.builder.build_report_workbook(
            outputs, out_path,
            customer="Apple", project="QAS-UFG",
            customer_po="PO", sales_order="SO",
            append_mode=False,
        )
        wb = openpyxl.load_workbook(out_path)
        summary = wb["Summary"]
        d_col_values = [
            summary.cell(r, 4).value for r in range(10, 13)
        ]
        non_empty = [v for v in d_col_values if v]
        self.assertEqual(
            len(non_empty), 2,
            f"Summary D column should list both devices, got "
            f"{d_col_values!r}",
        )
        # Both names carry their chassis serial.
        joined_d = " ".join(str(v) for v in non_empty)
        self.assertIn("M9A03562", joined_d)
        self.assertIn("M9A12345", joined_d)
        # F7 (device count) reflects two devices.
        self.assertEqual(summary["F7"].value, 2)

    def test_append_mode_serial_same_com_port_different_chassis(self):
        """The bug from the 2026-06-02 13:23 log: operator ran inventory
        on one WS5 via COM3, saved, then re-ran inventory in APPEND
        mode on a DIFFERENT WS5 via the same COM3. Both shelves report
        the post-provision hostname ``WS5_1``, the factory-default
        rewrite turns them into ``WS5-<serial>`` (unique per chassis),
        but the IP-fallback dedup heuristic was still firing — finding
        one prior device with the same IP key ``COM3`` and clobbering
        it. The fix gates IP-fallback on a flag that records whether
        the factory-default rewrite ran for the current device."""
        # First scan: build a workbook with one WS5.
        outputs_first = {
            "COM3": self._ws5_data("M9A03562", "M9A16485"),
        }
        out_path = os.path.join(self.tmp_dir, "append_test.xlsx")
        self.builder.build_report_workbook(
            outputs_first, out_path,
            customer="Apple", project="QAS-UFG",
            customer_po="PO", sales_order="SO",
            append_mode=False,
        )
        # Sanity: one device sheet exists.
        wb = openpyxl.load_workbook(out_path)
        first_pass = self._ws5_device_sheets(wb)
        self.assertEqual(first_pass, ["WS5_M9A03562"])
        wb.close()

        # Second scan: APPEND mode, same COM port, different chassis.
        outputs_second = {
            "COM3": self._ws5_data("M9A03390", "M9A2DCD6"),
        }
        self.builder.build_report_workbook(
            outputs_second, out_path,
            customer="Apple", project="QAS-UFG",
            customer_po="PO", sales_order="SO",
            append_mode=True,
        )

        # Both sheets must survive — neither chassis should be lost.
        wb = openpyxl.load_workbook(out_path)
        ws5_sheets = self._ws5_device_sheets(wb)
        self.assertEqual(
            ws5_sheets, ["WS5_M9A03390", "WS5_M9A03562"],
            f"Both WS5 chassis should keep distinct tabs after append; "
            f"got sheets {wb.sheetnames!r}",
        )

        summary = wb["Summary"]
        d_col_values = [
            summary.cell(r, 4).value for r in range(10, 13)
        ]
        non_empty = [v for v in d_col_values if v]
        self.assertEqual(
            len(non_empty), 2,
            f"Summary should list BOTH chassis after append, got "
            f"{d_col_values!r}",
        )
        self.assertEqual(summary["F7"].value, 2)


if __name__ == "__main__":
    unittest.main()
