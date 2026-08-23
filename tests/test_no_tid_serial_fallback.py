"""Tests for the no-TID → chassis-serial naming fallback.

Operator request: when a device reports no hostname / TID, use the
chassis serial as the device-tab name (and BoM column header, and
Summary row identifier) instead of the IP-based fallback. The latter
collapses multiple anonymous shelves behind the same management IP
(factory-default 10.0.0.1) onto a single shared tab.
"""
from __future__ import annotations
import os
import shutil
import tempfile
import unittest
from unittest.mock import MagicMock

import openpyxl
import pandas as pd

from gui.workbook_builder import (
    WorkbookBuilder,
    device_name_with_serial_fallback,
)


class TestDeviceNameWithSerialFallback(unittest.TestCase):
    """Pure helper tests — pin the fallback precedence."""

    def test_uses_primary_name_when_present(self):
        # The normal case: device reported a TID, just use it.
        self.assertEqual(
            device_name_with_serial_fallback(
                "lrt2.bb.net.apple.com", "M9A03562", "10.0.0.1",
            ),
            "lrt2.bb.net.apple.com",
        )

    def test_uses_serial_when_primary_name_empty(self):
        self.assertEqual(
            device_name_with_serial_fallback("", "M9A03562", "10.0.0.1"),
            "Device-M9A03562",
        )

    def test_uses_serial_when_primary_name_none(self):
        self.assertEqual(
            device_name_with_serial_fallback(None, "M9A03562", "10.0.0.1"),
            "Device-M9A03562",
        )

    def test_uses_serial_when_primary_name_nan(self):
        # pandas-side NaN values surface as the literal string 'nan'
        # after str() coercion — must be treated as empty.
        self.assertEqual(
            device_name_with_serial_fallback("nan", "M9A03562", "10.0.0.1"),
            "Device-M9A03562",
        )

    def test_falls_back_to_ip_when_no_name_or_serial(self):
        self.assertEqual(
            device_name_with_serial_fallback("", "", "10.0.0.1"),
            "System_10_0_0_1",
        )

    def test_falls_back_to_ip_when_name_and_serial_blank(self):
        self.assertEqual(
            device_name_with_serial_fallback("   ", "   ", "10.0.0.1"),
            "System_10_0_0_1",
        )

    def test_custom_prefixes_honored(self):
        # PSI builder passes ip_prefix="PSI" so anonymous PSI devices
        # surface as PSI_<ip> rather than System_<ip>.
        self.assertEqual(
            device_name_with_serial_fallback("", "", "172.16.0.1", ip_prefix="PSI"),
            "PSI_172_16_0_1",
        )
        self.assertEqual(
            device_name_with_serial_fallback("", "ABC123", "10.0.0.1",
                                             serial_prefix="Shelf"),
            "Shelf-ABC123",
        )

    def test_serial_with_whitespace_stripped(self):
        self.assertEqual(
            device_name_with_serial_fallback(
                "", "  M9A03562  ", "10.0.0.1",
            ),
            "Device-M9A03562",
        )


class TestBuildReportWorkbookUsesSerialFallback(unittest.TestCase):
    """End-to-end: ``build_report_workbook`` with two anonymous
    devices behind the same IP must produce two distinct device tabs,
    each named ``Device-<serial>``."""

    def setUp(self):
        self.tmp_dir = tempfile.mkdtemp(prefix="no_tid_fallback_")
        template_path = os.path.join(self.tmp_dir, "tpl.xlsx")
        tpl = openpyxl.Workbook()
        tpl.active.title = "Customer-project"
        tpl.create_sheet("Summary")
        tpl.create_sheet("DeviceTemplate")
        tpl.save(template_path)
        db = MagicMock()
        db.db_path = ":memory:"
        db.lookup_part.return_value = ""
        self.builder = WorkbookBuilder(
            db_cache=db, template_path=template_path,
            packing_slip_template="",
        )

    def tearDown(self):
        shutil.rmtree(self.tmp_dir, ignore_errors=True)

    @staticmethod
    def _anonymous_device_bundle(chassis_serial: str) -> dict:
        """Inventory-shape data WITHOUT a System Name — simulates a
        device that didn't report its TID. Chassis-row Serial Number
        carries *chassis_serial* so the fallback can pick it up."""
        df = pd.DataFrame([{
            # System Name intentionally blank to exercise the fallback.
            "System Name": "",
            "System Type": "Test System",
            "Type": "Shelf",
            "Information Type": "Shelf",
            "Part Number": "P-1",
            "Serial Number": chassis_serial,
            "Description": "Test chassis",
            "Name": "Main Shelf",
            "Source": "LAN",
        }])
        return {"main": df}

    def test_two_anonymous_devices_get_distinct_tabs_by_serial(self):
        out_path = os.path.join(self.tmp_dir, "out.xlsx")
        # Real-world scenario: two factory-default Ciena RLS shelves
        # both at 10.0.0.1 — but each chassis has its own serial.
        # Pass distinct keys at the outputs-dict layer (the upstream
        # collector already disambiguates, per the v2.0.6.1 work) so
        # both bundles reach the builder.
        self.builder.build_report_workbook(
            {
                "10.0.0.1": self._anonymous_device_bundle("M9A03562"),
                "10.0.0.1_dup": self._anonymous_device_bundle("M9A12345"),
            },
            out_path,
            customer="C", project="P",
            customer_po="PO", sales_order="SO",
            append_mode=False,
        )
        wb = openpyxl.load_workbook(out_path)
        # Sheet titles sanitize the "-" in Device-<serial> to "_"
        # via make_unique_sheet_title.
        device_tabs = sorted(
            n for n in wb.sheetnames if n.startswith("Device_")
        )
        self.assertEqual(
            device_tabs, ["Device_M9A03562", "Device_M9A12345"],
            f"Two anonymous shelves should get distinct "
            f"Device-<serial> tabs; got sheets={wb.sheetnames!r}",
        )

    def test_named_device_unchanged_by_fallback(self):
        # The fallback must NOT affect devices that DID report a TID.
        out_path = os.path.join(self.tmp_dir, "named.xlsx")
        df = pd.DataFrame([{
            "System Name": "real-hostname-1",
            "System Type": "Test", "Type": "Shelf",
            "Information Type": "Shelf",
            "Part Number": "P", "Serial Number": "SHOULDNT-MATTER",
            "Description": "d", "Name": "Main", "Source": "LAN",
        }])
        self.builder.build_report_workbook(
            {"10.0.0.1": {"main": df}}, out_path,
            customer="C", project="P",
            customer_po="PO", sales_order="SO",
            append_mode=False,
        )
        wb = openpyxl.load_workbook(out_path)
        # make_unique_sheet_title sanitizes "-" to "_" in the title.
        # The reported hostname survives in F6 of the device tab (the
        # un-sanitized form) — that's what we actually check.
        self.assertIn("real_hostname_1", wb.sheetnames)
        self.assertEqual(
            wb["real_hostname_1"]["F6"].value, "real-hostname-1",
        )
        # No Device-<serial> tab created when a name was reported.
        device_tabs = [n for n in wb.sheetnames if n.startswith("Device_")]
        self.assertEqual(device_tabs, [])

    def test_truly_anonymous_falls_back_to_ip_name(self):
        # No TID AND no chassis serial → IP-based name (last resort).
        out_path = os.path.join(self.tmp_dir, "anon.xlsx")
        df = pd.DataFrame([{
            "System Name": "", "System Type": "Test",
            "Type": "Shelf", "Information Type": "Shelf",
            "Part Number": "P", "Serial Number": "",  # no serial either
            "Description": "d", "Name": "Main", "Source": "LAN",
        }])
        self.builder.build_report_workbook(
            {"10.0.0.5": {"main": df}}, out_path,
            customer="C", project="P",
            customer_po="PO", sales_order="SO",
            append_mode=False,
        )
        wb = openpyxl.load_workbook(out_path)
        # System_10_0_0_5 fallback fires; sheet title sanitizes the
        # leading 'S' through 'S' (no chars to swap).
        self.assertIn("System_10_0_0_5", wb.sheetnames)


if __name__ == "__main__":
    unittest.main()
