"""Tests for the Asset Import parser + merger (utils/asset_import.py)."""
from __future__ import annotations
import os
import shutil
import tempfile
import unittest

import openpyxl

from utils.asset_import import (
    AssetImportResult,
    AssetRecord,
    apply_asset_data_to_inventory,
    parse_asset_doc,
    _find_header_row,
    _norm,
    _norm_serial,
)


class TestNormHelpers(unittest.TestCase):

    def test_header_normalization_collapses_whitespace_punct(self):
        # Real-world variants we must accept as the same header.
        self.assertEqual(_norm("PO"), "po")
        self.assertEqual(_norm("po"), "po")
        self.assertEqual(_norm("PO Number"), "ponumber")
        self.assertEqual(_norm("Serial #"), "serial")
        self.assertEqual(_norm("Serial Number"), "serialnumber")
        self.assertEqual(_norm("Asset"), "asset")
        self.assertEqual(_norm("Asset #"), "asset")
        self.assertEqual(_norm(None), "")
        self.assertEqual(_norm(""), "")

    def test_serial_normalization_uppercases_and_strips(self):
        self.assertEqual(_norm_serial("nntmrt1l3k6n"), "NNTMRT1L3K6N")
        self.assertEqual(_norm_serial("  NN-12345  "), "NN-12345")
        self.assertEqual(_norm_serial(None), "")
        self.assertEqual(_norm_serial("nan"), "")
        self.assertEqual(_norm_serial(""), "")


def _write_asn_workbook(
    path: str,
    *,
    header_row: int = 22,
    headers: list = None,
    data_rows: list = None,
    metadata_rows: int = 18,
) -> None:
    """Build a minimal ASN-shaped workbook for parser tests.

    *header_row* / *headers* control the header position + labels.
    *data_rows* is a list of dicts keyed by header name. Anything
    before the header row is filled with "Customer:" / "PO Number:"
    style metadata so the parser has to skip past it.
    """
    if headers is None:
        headers = ["Line #", "Qty", "PO", "Vendor", "Customer Model",
                   "Description", "Serial #", "Asset", "Parent Asset"]
    if data_rows is None:
        data_rows = []
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ASN"
    # Header preamble
    ws["A1"] = "Advanced Ship Notification"
    ws["A3"] = "Customer:"
    ws["A5"] = "PO Number:"
    ws["B5"] = "6000759143"
    # Headers on the configured row
    for col_idx, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col_idx, value=h)
    # Data rows
    for row_offset, rec in enumerate(data_rows):
        r = header_row + 1 + row_offset
        for col_idx, h in enumerate(headers, start=1):
            if h in rec:
                ws.cell(row=r, column=col_idx, value=rec[h])
    wb.save(path)


class TestParseAssetDoc(unittest.TestCase):

    def setUp(self):
        self.tmp_dir = tempfile.mkdtemp(prefix="asset_import_")

    def tearDown(self):
        shutil.rmtree(self.tmp_dir, ignore_errors=True)

    def _path(self, name: str = "asn.xlsx") -> str:
        return os.path.join(self.tmp_dir, name)

    def test_parses_real_world_asn_shape(self):
        """The user's example file has headers at row 22 with columns
        Line #, Qty Ordered, Qty Shipped, PO, Vendor Model #, Customer
        Model #, Description, Serial #, Asset, Parent Asset. Reproduce
        that exact layout and assert the parser pulls the right map."""
        p = self._path()
        _write_asn_workbook(
            p,
            header_row=22,
            headers=[
                "Line #", "Qty Ordered", "Qty Shipped", "PO",
                "Vendor Model #", "Customer Model #", "Description",
                "Serial #", "Asset", "Parent Asset",
            ],
            data_rows=[
                {
                    "PO": 6000759143,
                    "Serial #": "NNTMRT1L3K6N",
                    "Asset": "A3535551",
                },
                {
                    "PO": 6000759143,
                    "Serial #": "NNTMRT1L18YH",
                    "Asset": "A3535552",
                },
            ],
        )
        result = parse_asset_doc(p)
        self.assertEqual(len(result), 2)
        self.assertEqual(
            result["NNTMRT1L3K6N"],
            AssetRecord(po="6000759143", asset="A3535551"),
        )
        self.assertEqual(
            result["NNTMRT1L18YH"],
            AssetRecord(po="6000759143", asset="A3535552"),
        )

    def test_finds_header_regardless_of_row(self):
        # Some ASN templates put headers higher / lower. Header
        # detection must scan, not hard-code a row.
        for hr in (3, 11, 22, 45):
            p = self._path(f"hr{hr}.xlsx")
            _write_asn_workbook(
                p, header_row=hr,
                data_rows=[{
                    "PO": "PO-X", "Serial #": "SER-X", "Asset": "AST-X",
                }],
            )
            result = parse_asset_doc(p)
            self.assertEqual(
                result, {"SER-X": AssetRecord(po="PO-X", asset="AST-X")},
                f"header at row {hr} failed",
            )

    def test_lowercases_serial_normalized_to_upper(self):
        p = self._path()
        _write_asn_workbook(
            p,
            data_rows=[{
                "PO": "PO-1", "Serial #": "  nn-lowercase  ",
                "Asset": "A-1",
            }],
        )
        result = parse_asset_doc(p)
        self.assertIn("NN-LOWERCASE", result)
        self.assertNotIn("nn-lowercase", result)

    def test_skips_rows_without_serial(self):
        p = self._path()
        _write_asn_workbook(
            p,
            data_rows=[
                {"PO": "PO-1", "Asset": "A-1"},     # missing Serial
                {"PO": "PO-1", "Serial #": "S-2",
                 "Asset": "A-2"},
            ],
        )
        result = parse_asset_doc(p)
        self.assertEqual(set(result.keys()), {"S-2"})

    def test_skips_rows_with_no_po_and_no_asset(self):
        p = self._path()
        _write_asn_workbook(
            p,
            data_rows=[
                {"Serial #": "S-EMPTY"},  # just a serial; useless
                {"PO": "PO-1", "Serial #": "S-OK", "Asset": "A-1"},
            ],
        )
        result = parse_asset_doc(p)
        self.assertEqual(set(result.keys()), {"S-OK"})

    def test_duplicate_serial_last_wins(self):
        p = self._path()
        _write_asn_workbook(
            p,
            data_rows=[
                {"PO": "PO-1", "Serial #": "DUP", "Asset": "A-1"},
                {"PO": "PO-2", "Serial #": "DUP", "Asset": "A-2"},
            ],
        )
        result = parse_asset_doc(p)
        self.assertEqual(result["DUP"], AssetRecord(po="PO-2", asset="A-2"))

    def test_missing_headers_raises_value_error(self):
        # ASN with only PO + Serial (no Asset column) is unparseable.
        p = self._path()
        _write_asn_workbook(
            p, headers=["Line #", "PO", "Serial #", "Description"],
            data_rows=[{"PO": "X", "Serial #": "S-1"}],
        )
        with self.assertRaises(ValueError):
            parse_asset_doc(p)

    def test_serial_pound_variants_all_match(self):
        # Header text variants the parser must recognize as "Serial #".
        for label in ("Serial #", "Serial Number", "serial-number", "SERIAL#"):
            p = self._path(f"{_norm(label)}.xlsx")
            _write_asn_workbook(
                p,
                headers=["PO", label, "Asset"],
                data_rows=[{"PO": "P", label: "S-X", "Asset": "A-X"}],
            )
            self.assertIn("S-X", parse_asset_doc(p), f"label {label!r} failed")


# ── Merger tests ──────────────────────────────────────────────────────


def _add_device_tab(
    wb,
    name: str,
    hostname: str,
    rows: list,
    pre_existing_po: str = "",
    pre_existing_asset_at_g15: str = "",
):
    """Create a device tab in the ATLAS layout (F6=hostname,
    C7=Customer PO if pre-set, data row 15+ with B/C/D/E/F populated
    and optional G).
    """
    ws = wb.create_sheet(name)
    ws["F6"] = hostname
    if pre_existing_po:
        ws["C7"] = pre_existing_po
    for i, rec in enumerate(rows):
        r = 15 + i
        ws.cell(row=r, column=2, value=rec.get("name", ""))
        ws.cell(row=r, column=3, value=rec.get("type", ""))
        ws.cell(row=r, column=4, value=rec.get("part", ""))
        ws.cell(row=r, column=5, value=rec.get("serial", ""))
        ws.cell(row=r, column=6, value=rec.get("desc", ""))
        if "asset" in rec:
            ws.cell(row=r, column=7, value=rec["asset"])
    if pre_existing_asset_at_g15:
        ws["G15"] = pre_existing_asset_at_g15


class TestApplyAssetDataToInventory(unittest.TestCase):

    def _bare_workbook(self):
        wb = openpyxl.Workbook()
        # Drop the default sheet so tests start with a clean slate.
        wb.remove(wb.active)
        wb.create_sheet("Summary")
        return wb

    def test_writes_asset_and_po_for_matching_serials(self):
        wb = self._bare_workbook()
        _add_device_tab(
            wb, "dev_one", "device-one.example.com",
            rows=[
                {"name": "Shelf", "type": "Shelf", "serial": "S-CHASSIS"},
                {"name": "Card1", "type": "Card", "serial": "S-CARD1"},
            ],
        )
        asset_map = {
            "S-CHASSIS": AssetRecord(po="PO-100", asset="A-CHASSIS"),
            "S-CARD1": AssetRecord(po="PO-100", asset="A-CARD1"),
        }
        result = apply_asset_data_to_inventory(wb, asset_map)

        ws = wb["dev_one"]
        self.assertEqual(ws["G15"].value, "A-CHASSIS")
        self.assertEqual(ws["G16"].value, "A-CARD1")
        self.assertEqual(ws["C7"].value, "PO-100")
        self.assertEqual(result.tabs_touched, 1)
        self.assertEqual(result.rows_matched, 2)
        self.assertEqual(result.tabs_with_no_matches, [])
        self.assertEqual(result.po_conflicts, [])

    def test_overwrites_xlookup_formula_at_g15(self):
        # The fresh inventory build plants =XLOOKUP(...) at G15.
        # Asset Import is bulk-fill — it must replace the formula with
        # the imported literal value.
        wb = self._bare_workbook()
        _add_device_tab(
            wb, "dev_a", "host-a",
            rows=[{"serial": "S-1"}],
            pre_existing_asset_at_g15=(
                '=XLOOKUP(F6,Summary!$D:$D,Summary!$E:$E,"",0,1)'
            ),
        )
        asset_map = {"S-1": AssetRecord(po="PO-1", asset="A-1")}
        apply_asset_data_to_inventory(wb, asset_map)
        ws = wb["dev_a"]
        self.assertEqual(ws["G15"].value, "A-1")
        self.assertFalse(
            str(ws["G15"].value).startswith("="),
            "Formula must be overwritten with the imported literal",
        )

    def test_serial_match_is_case_insensitive(self):
        wb = self._bare_workbook()
        _add_device_tab(
            wb, "dev_a", "host-a",
            rows=[{"serial": "nn-low-case"}],
        )
        asset_map = {"NN-LOW-CASE": AssetRecord(po="PO-X", asset="A-X")}
        apply_asset_data_to_inventory(wb, asset_map)
        self.assertEqual(wb["dev_a"]["G15"].value, "A-X")

    def test_skips_summary_and_bom_sheets(self):
        wb = self._bare_workbook()
        wb.create_sheet("BOM")
        # Plant a fake serial in both Summary and BOM — neither
        # should be considered a device tab.
        wb["Summary"]["E15"] = "S-X"
        wb["BOM"]["E15"] = "S-X"
        asset_map = {"S-X": AssetRecord(po="PO-1", asset="A-1")}
        result = apply_asset_data_to_inventory(wb, asset_map)
        self.assertEqual(result.tabs_touched, 0)
        self.assertNotEqual(wb["Summary"]["G15"].value, "A-1")
        self.assertNotEqual(wb["BOM"]["G15"].value, "A-1")

    def test_po_conflict_leaves_c7_untouched_and_flags(self):
        wb = self._bare_workbook()
        _add_device_tab(
            wb, "dev_mixed", "host-mixed",
            rows=[
                {"serial": "S-FROM-PO-1"},
                {"serial": "S-FROM-PO-2"},
            ],
            pre_existing_po="ORIGINAL-PO",
        )
        asset_map = {
            "S-FROM-PO-1": AssetRecord(po="PO-1", asset="A-1"),
            "S-FROM-PO-2": AssetRecord(po="PO-2", asset="A-2"),
        }
        result = apply_asset_data_to_inventory(wb, asset_map)
        ws = wb["dev_mixed"]
        # Both assets still got written (matching rows are fine).
        self.assertEqual(ws["G15"].value, "A-1")
        self.assertEqual(ws["G16"].value, "A-2")
        # C7 preserved — we refuse to pick arbitrarily on conflict.
        self.assertEqual(ws["C7"].value, "ORIGINAL-PO")
        self.assertEqual(len(result.po_conflicts), 1)
        conflict_tab, pos = result.po_conflicts[0]
        self.assertEqual(conflict_tab, "dev_mixed")
        self.assertEqual(sorted(pos), ["PO-1", "PO-2"])

    def test_tab_with_no_matches_reported(self):
        wb = self._bare_workbook()
        _add_device_tab(
            wb, "dev_unknown", "host-unknown",
            rows=[{"serial": "NOT-IN-ASSET-DOC"}],
        )
        result = apply_asset_data_to_inventory(
            wb, {"OTHER-SERIAL": AssetRecord(po="P", asset="A")},
        )
        self.assertEqual(result.tabs_touched, 0)
        self.assertIn("dev_unknown", result.tabs_with_no_matches)
        self.assertEqual(result.rows_matched, 0)
        # The unmatched serial got counted.
        self.assertGreaterEqual(result.serials_not_in_asset_doc, 1)

    def test_blank_asset_record_still_stamps_po(self):
        # An ASN row that has a PO but no Asset (rare but possible)
        # should still propagate the PO to C7 for matched serials,
        # without writing anything to G.
        wb = self._bare_workbook()
        _add_device_tab(
            wb, "dev_x", "host-x",
            rows=[{"serial": "S-NO-ASSET"}],
        )
        asset_map = {"S-NO-ASSET": AssetRecord(po="PO-99", asset="")}
        apply_asset_data_to_inventory(wb, asset_map)
        ws = wb["dev_x"]
        self.assertEqual(ws["C7"].value, "PO-99")
        # G15 untouched (no asset to write).
        self.assertIsNone(ws["G15"].value)


class TestFindHeaderRow(unittest.TestCase):
    """Direct guards for the header-detection helper used by the parser."""

    def test_requires_all_three_columns_on_same_row(self):
        # A row with PO + Serial but missing Asset must NOT count as
        # the header — _find_header_row returns None.
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "PO"
        ws["B1"] = "Serial #"
        # No Asset on row 1.
        ws["A2"] = "PO Number"  # but no Serial / Asset here either
        self.assertIsNone(_find_header_row(ws))

    def test_returns_one_based_column_positions(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["B3"] = "PO"
        ws["D3"] = "Serial #"
        ws["F3"] = "Asset"
        anchor = _find_header_row(ws)
        self.assertEqual(anchor, (2, 4, 6, 3))


if __name__ == "__main__":
    unittest.main()
