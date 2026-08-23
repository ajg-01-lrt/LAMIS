"""Regression tests for the Summary-to-tab hyperlink resolution path.

Beta-test artifact ``PostFileProcessing_BomProblem.BOM.xlsx`` exposed a
silent failure: every BOM Total dropped to 0 and the per-site column was
gone. Root cause was the hyperlink-target regex.

``_populate_summary_table`` writes Summary D-column hyperlinks as
``#'<sheet_title>'!A1`` (note the leading ``#``). The original
``_HYPERLINK_TARGET_RE`` only accepted forms WITHOUT the ``#``, so
``_resolve_hyperlink_target`` returned ``None`` for every Summary row in
real BOM workbooks. ``_extract_summary`` then had no way to map the
display label (dotted FQDN form) back to the underscored sheet tab, and
returned an empty list — which propagated to ``_build_bom_sheet`` as
"no sites" and zeroed every Total.
"""
from __future__ import annotations
import unittest
from unittest.mock import MagicMock

import openpyxl
from openpyxl.worksheet.hyperlink import Hyperlink

from gui import bom_frame
from gui.bom_frame import _HYPERLINK_TARGET_RE, _resolve_hyperlink_target, BomFrame


class TestHyperlinkTargetRegex(unittest.TestCase):
    """Targets the bare regex so future edits can't silently regress."""

    def test_hash_quoted_sheet_form_matches(self):
        # Form actually produced by _populate_summary_table:
        m = _HYPERLINK_TARGET_RE.match("#'usufg1_l8r1_bb_net_apple_com'!A1")
        self.assertIsNotNone(m, "must accept the #'sheet'!cell form")
        self.assertEqual(m.group(1), "usufg1_l8r1_bb_net_apple_com")

    def test_quoted_sheet_form_still_matches(self):
        m = _HYPERLINK_TARGET_RE.match("'usufg1_l8r1_bb_net_apple_com'!A1")
        self.assertIsNotNone(m)
        self.assertEqual(m.group(1), "usufg1_l8r1_bb_net_apple_com")

    def test_plain_sheet_form_still_matches(self):
        m = _HYPERLINK_TARGET_RE.match("Summary!A1")
        self.assertIsNotNone(m)
        self.assertEqual(m.group(1), "Summary")


class TestResolveHyperlinkTarget(unittest.TestCase):
    """End-to-end on an openpyxl cell, since that's how the code reads
    workbooks in production. Matches the form Excel actually stores."""

    def test_resolves_workbook_written_hyperlink(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Summary"
        wb.create_sheet("usufg1_l8r1_bb_net_apple_com")
        # Replicate _populate_summary_table line 463 exactly:
        ws["D10"].value = "usufg1-l8r1.bb.net.apple.com"
        ws["D10"].hyperlink = "#'usufg1_l8r1_bb_net_apple_com'!A1"
        ws["D10"].style = "Hyperlink"

        resolved = _resolve_hyperlink_target(ws["D10"])
        self.assertEqual(resolved, "usufg1_l8r1_bb_net_apple_com")

    def test_returns_none_for_cell_without_hyperlink(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"].value = "no link"
        self.assertIsNone(_resolve_hyperlink_target(ws["A1"]))


class TestExtractSummaryWithDottedLabels(unittest.TestCase):
    """Bug-of-the-day: Summary D10 carries the dotted display form
    (``usufg1-l8r1.bb.net.apple.com``) while the sheet tab carries the
    underscored form. ``_extract_summary`` must reconcile via hyperlink
    or the BOM rebuild produces empty site columns."""

    def _make_workbook(self) -> openpyxl.Workbook:
        wb = openpyxl.Workbook()
        summary = wb.active
        summary.title = "Summary"
        # Two devices, both with dotted FQDN labels but underscored sheets.
        for sheet in ("usufg1_l8r1_bb_net_apple_com", "usufg1_l8r2_bb_net_apple_com"):
            wb.create_sheet(sheet)
        summary["B9"], summary["C9"], summary["D9"] = "#", "IP Address", "Device Name"

        summary["B10"], summary["C10"] = 1, "10.0.0.1"
        summary["D10"].value = "usufg1-l8r1.bb.net.apple.com"
        summary["D10"].hyperlink = "#'usufg1_l8r1_bb_net_apple_com'!A1"

        summary["B11"], summary["C11"] = 2, "10.0.0.5"
        summary["D11"].value = "usufg1-l8r2.bb.net.apple.com"
        summary["D11"].hyperlink = "#'usufg1_l8r2_bb_net_apple_com'!A1"
        return wb

    def test_extract_finds_both_devices_via_hyperlink(self):
        wb = self._make_workbook()
        ws = wb["Summary"]
        items, display_to_tab = BomFrame._extract_summary(ws, set(wb.sheetnames))
        # Two devices, dotted display, underscored tab.
        self.assertEqual(len(items), 2, f"expected 2 items, got {items!r}")
        labels = {item[1] for item in items}
        self.assertIn("usufg1-l8r1.bb.net.apple.com", labels)
        self.assertIn("usufg1-l8r2.bb.net.apple.com", labels)
        # Mapping is dotted-display -> underscored-tab.
        self.assertEqual(
            display_to_tab["usufg1-l8r1.bb.net.apple.com"],
            "usufg1_l8r1_bb_net_apple_com",
        )

    def test_extract_returns_empty_when_no_hyperlink_and_value_not_sheetname(self):
        """Sanity: if hyperlink resolution genuinely fails (no link AND the
        value doesn't match a sheetname), the row gets skipped — that's
        the original behavior, which the fix must preserve."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Summary"
        wb.create_sheet("device_a")
        ws["D10"].value = "totally-unrelated-label"  # no hyperlink, no sheet match
        items, _ = BomFrame._extract_summary(ws, set(wb.sheetnames))
        self.assertEqual(items, [])


if __name__ == "__main__":
    unittest.main()
