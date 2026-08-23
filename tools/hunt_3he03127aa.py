"""Hunt for 3HE03127AA across the Sales BoM AND Factory Live Inventory —
is it really a surplus, or do we have a matching PN that's not aliased?"""
from __future__ import annotations
import re
from pathlib import Path
from openpyxl import load_workbook

SALES_SRC = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/ATLAS Build Material/TO4-Services-Prices-with-BoM-V9.1-03-05-26_BPA.xlsx")
SALES_OUT = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/Sales_BPA_NGT_Refresh_TO4_SalesBoM_Site_Packing_2026-05-23_00-33.xlsx")
LIVE      = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/Master_BPA_TO4_Factoy_Inventory.BOM.xlsx")

# Substrings to flag — exact + relaxed forms
NEEDLES = ("3HE03127", "03127AA", "OC3/STM1 CHANNELIZED", "CHANNELIZED ADAPTER")

def norm(s): return re.sub(r"\s+", "", str(s or "")).upper()

def scan(path, label):
    print(f"\n{'='*72}\n{label}: {path.name}\n{'='*72}")
    if not path.exists():
        print("  (missing)"); return
    wb = load_workbook(path, data_only=True)
    print(f"  Sheets: {len(wb.sheetnames)}")
    hits = 0
    for ws in wb.worksheets:
        max_r = min(ws.max_row or 0, 5000)
        max_c = min(ws.max_column or 0, 200)
        for row in ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c, values_only=False):
            for cell in row:
                if cell.value is None: continue
                sv_upper = str(cell.value).upper()
                if any(n.upper() in sv_upper for n in NEEDLES):
                    # show only PN-like cells (in col B 2 or col D 4) or desc col
                    if cell.column in (2, 3, 4, 6):
                        rv = [(c, ws.cell(cell.row, c).value)
                              for c in range(1, max_c+1)]
                        rv = [(c, v) for c, v in rv if v not in (None, "")][:18]
                        print(f"  [{ws.title}] {cell.coordinate} value={cell.value!r}")
                        print(f"     row: {rv}")
                        hits += 1
                        if hits >= 25: return

scan(SALES_SRC, "SALES SOURCE")
scan(SALES_OUT, "SALES OUTPUT")
scan(LIVE,      "FACTORY LIVE")
