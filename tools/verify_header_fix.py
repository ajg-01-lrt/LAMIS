"""Smoke-test that _parse_per_site_bom now accepts the new Sales BoM."""
import sys
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from gui.bom_compare_frame import BomCompareFrame

SALES = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/Sales_BPA_NGT_Refresh_TO4_SalesBoM_Site_Packing_2026-05-23_00-33.xlsx")
LIVE  = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/Master_BPA_TO4_Factoy_Inventory.BOM.xlsx")

wb_sales = load_workbook(SALES, data_only=True)
wb_live  = load_workbook(LIVE, data_only=True)
ws_sales = wb_sales["BOM"]
ws_live  = wb_live["BOM"]

# create a bare instance just to call the bound method (it's a regular method
# that takes `self` but never touches Tk state)
frame = BomCompareFrame.__new__(BomCompareFrame)

try:
    site_cols, per_site, spares, parts = frame._parse_per_site_bom(ws_sales)
    print(f"SALES OK: {len(site_cols)} site cols, {len(per_site)} per-site PNs, "
          f"{len(spares)} spares PNs, {len(parts)} unique parts")
    print(f"  first 5 sites: {site_cols[:5]}")
    sample = list(per_site.items())[:3]
    print(f"  sample per-site (3): {sample}")
except Exception as e:
    print(f"SALES FAILED: {e!r}")

try:
    site_cols, per_site, spares, parts = frame._parse_per_site_bom(ws_live)
    print(f"\nLIVE OK: {len(site_cols)} site cols, {len(per_site)} per-site PNs, "
          f"{len(spares)} spares PNs, {len(parts)} unique parts")
    print(f"  first 5 sites: {site_cols[:5]}")
except Exception as e:
    print(f"\nLIVE FAILED: {e!r}")
