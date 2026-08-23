from openpyxl import load_workbook
from pathlib import Path
SRC = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/ATLAS Build Material/TO4-Services-Prices-with-BoM-V9.1-03-05-26_BPA.xlsx")
wb = load_workbook(SRC, data_only=True)
ws = wb["TO4 BoM v8"]
print(f"max_col={ws.max_column}")
for c in range(1, ws.max_column+1):
    r6 = ws.cell(6, c).value
    r7 = ws.cell(7, c).value
    r10 = ws.cell(10, c).value
    r11 = ws.cell(11, c).value
    if any(v is not None for v in (r6, r7, r10, r11)):
        print(f"  col {c:>3}: r6={r6!r:30} r7={r7!r:30} r10={r10!r:20} r11={r11!r}")
