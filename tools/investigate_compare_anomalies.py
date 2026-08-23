"""Drill into the new Sales BoM to verify suspicious COMPARE rows."""
from __future__ import annotations
import re
from pathlib import Path
from openpyxl import load_workbook

SALES = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/Sales_BPA_NGT_Refresh_TO4_SalesBoM_Site_Packing_2026-05-23_00-33.xlsx")
LIVE  = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/Master_BPA_TO4_Factoy_Inventory.BOM.xlsx")

def norm(s): return re.sub(r"\s+", "", str(s or "")).upper()

# 1. Does the new Sales BoM actually contain NMA-8509 anywhere?
print("=" * 70)
print("1. NMA-8509 search across new Sales BoM (full workbook)")
print("=" * 70)
wb = load_workbook(SALES, data_only=True)
for ws in wb.worksheets:
    max_r, max_c = ws.max_row or 0, ws.max_column or 0
    if max_r > 5000 or max_c > 200:
        max_r = min(max_r, 5000); max_c = min(max_c, 200)
    for row in ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c, values_only=False):
        for cell in row:
            if cell.value is None: continue
            sv = str(cell.value)
            if "NMA-8509" in sv or "8509" in sv:
                # show full row
                rv = [(c, ws.cell(cell.row, c).value) for c in range(1, max_c+1)]
                rv = [(c, v) for c, v in rv if v not in (None, "")]
                print(f"  [{ws.title}] {cell.coordinate} value={sv!r}")
                print(f"     row {cell.row} cells: {rv[:30]}")

# 2. The new Sales BoM — does 3HE04824AA actually appear now?
print()
print("=" * 70)
print("2. 3HE04824AA across new Sales BoM (was 0 in prior run, now showing 92)")
print("=" * 70)
target = norm("3HE04824AA")
for ws in wb.worksheets:
    max_r = min(ws.max_row or 0, 2000)
    max_c = min(ws.max_column or 0, 100)
    for row in ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c, values_only=False):
        for cell in row:
            if cell.value is None: continue
            if norm(cell.value) == target:
                rv = [(c, ws.cell(cell.row, c).value) for c in range(1, max_c+1)]
                rv = [(c, v) for c, v in rv if v not in (None, "")]
                print(f"  [{ws.title}] {cell.coordinate} row {cell.row}: {rv[:40]}")

# 3. The new Sales BoM BOM tab — get row 7 (headers) and the row for 3HE04824AA
print()
print("=" * 70)
print("3. New Sales BoM aggregate BoM — site columns ranking by SAR-8 row")
print("=" * 70)
ws = wb["BOM"]
# Reading 3HE04824AA row + 3HE04823AA row
for tgt in ("3HE04824AA", "3HE04823AA"):
    tn = norm(tgt)
    for r in range(10, (ws.max_row or 10)+1):
        pn = ws.cell(r, 2).value
        if pn and norm(pn) == tn:
            print(f"\n  PN={tgt} on row {r}")
            for c in range(1, (ws.max_column or 0)+1):
                v = ws.cell(r, c).value
                if v not in (None, "", 0):
                    h = ws.cell(7, c).value
                    print(f"     col {c:>3} ({h!r}): {v!r}")
            break
