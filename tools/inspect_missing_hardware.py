"""Look at the suspicious hardware-looking PNs that have qty ONLY in
col 67 (Network) or col 72 (License points) on TO4 BoM v8. Show the
whole row so we can decide if they're real orders."""
from openpyxl import load_workbook
from pathlib import Path
SRC = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/ATLAS Build Material/TO4-Services-Prices-with-BoM-V9.1-03-05-26_BPA.xlsx")
wb = load_workbook(SRC, data_only=True)
ws = wb["TO4 BoM v8"]

TARGETS = ("3HE04962AB", "3HE06151AC", "3HE07942AA", "3HE11280AA",
           "3HE12518AA", "3HE14004AA", "3HE16713AA", "3HE16716AA", "3HE16718AA")

# First print the relevant header rows for cols 65-80
print("HEADERS (rows 9-12, cols 65-80):")
for r in (9, 10, 11, 12):
    cells = []
    for c in range(65, 81):
        v = ws.cell(r, c).value
        if v is not None:
            cells.append(f"col{c}={v!r}")
    if cells:
        print(f"  r{r}: " + " | ".join(cells))

print("\nDATA rows for missing hardware PNs:")
for r in range(12, ws.max_row+1):
    pn = ws.cell(r, 2).value
    if pn and str(pn).strip().upper() in TARGETS:
        desc = ws.cell(r, 3).value
        print(f"\n  r{r} PN={pn!r} desc={desc!r}")
        nonzero = []
        for c in range(1, ws.max_column+1):
            v = ws.cell(r, c).value
            if v not in (None, "", 0):
                hdr10 = ws.cell(10, c).value or ""
                hdr11 = ws.cell(11, c).value or ""
                hdr_label = f"{hdr10}|{hdr11}".strip("|")
                nonzero.append((c, hdr_label, v))
        for c, hdr, v in nonzero:
            print(f"    col {c:>3} [{hdr!s:25}]: {v!r}")
