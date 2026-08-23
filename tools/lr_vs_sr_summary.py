"""Cross-tab: how many SR vs LR 10GE SFPs ordered (Sales) vs on-hand (Live)?"""
from __future__ import annotations
import re
from pathlib import Path
from openpyxl import load_workbook

LIVE_PATH    = Path(r"c:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/Master_BPA_TO4_Factoy_Inventory.BOM.xlsx")
SALES_PATH   = Path(r"c:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/TO4_BPA_NGT_Refresh_SalesBoM_Site_Packing_2026-05-22_22-05.xlsx")

TARGETS = ["3HE04823AA", "3HE04824AA"]

def norm(s): return re.sub(r"\s+","",str(s or "")).upper()

def count_live_per_site(target):
    wb = load_workbook(LIVE_PATH, data_only=True)
    site_counts = {}
    target_n = norm(target)
    for ws in wb.worksheets:
        if ws.title.lower() in ("bom","master") or "summary" in ws.title.lower():
            continue
        cnt = 0
        max_r = min(ws.max_row or 0, 5000)
        for r in range(1, max_r+1):
            v = ws.cell(row=r, column=4).value  # col D = PN on per-site tabs
            if v and norm(v) == target_n:
                cnt += 1
        if cnt:
            site_counts[ws.title] = cnt
    return site_counts

def count_sales_per_site(target):
    wb = load_workbook(SALES_PATH, data_only=True)
    target_n = norm(target)
    # In Sales BoM, the BoM aggregate sheet has per-site columns. But simpler: scan per-site tabs.
    site_counts = {}
    for ws in wb.worksheets:
        if ws.title.lower() in ("summary","bom","spares"):
            continue
        # find PN col (likely col B = 2 from earlier)
        max_r = min(ws.max_row or 0, 1000)
        cnt = 0
        for r in range(15, max_r+1):
            pn = ws.cell(row=r, column=2).value
            if pn and norm(pn) == target_n:
                # find quantity columns - typically col D onward
                qty = 0
                for c in range(4, min(ws.max_column or 4, 30)+1):
                    v = ws.cell(row=r, column=c).value
                    try:
                        qty += int(v)
                    except Exception:
                        pass
                if qty:
                    cnt += qty
        if cnt:
            site_counts[ws.title] = cnt
    return site_counts

for tgt in TARGETS:
    print(f"\n========== {tgt} ==========")
    live = count_live_per_site(tgt)
    sales = count_sales_per_site(tgt)
    print(f"  LIVE on-hand: total={sum(live.values())}  per site: {live}")
    print(f"  SALES ordered: total={sum(sales.values())}  per site: {sales}")
