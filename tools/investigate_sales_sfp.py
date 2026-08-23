"""List every PN+desc pair in the Sales BoM that looks SFP-ish, so we can see
what 10GE optic was actually ordered (if any)."""
from __future__ import annotations
import re
from pathlib import Path
from openpyxl import load_workbook

SALES_PATH = Path(r"c:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/TO4_BPA_NGT_Refresh_SalesBoM_Site_Packing_2026-05-22_22-05.xlsx")

KEYWORDS = ("SFP", "10GE", "10G", "OPTIC", "TRANSCEIVER", "XCVR")

def main():
    wb = load_workbook(SALES_PATH, data_only=True)
    print(f"Sheets: {wb.sheetnames}")
    for ws in wb.worksheets:
        print(f"\n=== {ws.title} ===")
        max_r = min(ws.max_row or 0, 2000)
        max_c = min(ws.max_column or 0, 60)
        # find header row
        hdr_row = None
        hdrs = []
        for r in range(1, min(20, max_r + 1)):
            vals = [str(ws.cell(row=r, column=c).value or "") for c in range(1, max_c + 1)]
            joined = " | ".join(vals).lower()
            if ("part" in joined or "material" in joined) and ("description" in joined or "desc" in joined):
                hdr_row = r
                hdrs = vals
                break
        if not hdr_row:
            print("  (no header row)")
            continue
        # locate columns
        pn_col = desc_col = total_col = spares_col = None
        for i, h in enumerate(hdrs, start=1):
            hl = h.lower().strip()
            if pn_col is None and ("part" in hl and ("number" in hl or "no" in hl or "#" in hl) or hl == "material"):
                pn_col = i
            if desc_col is None and ("description" in hl or hl == "desc"):
                desc_col = i
            if total_col is None and hl == "total":
                total_col = i
            if spares_col is None and "spare" in hl:
                spares_col = i
        print(f"  header row {hdr_row}: PN col={pn_col}, Desc col={desc_col}, Total col={total_col}, Spares col={spares_col}")
        for r in range(hdr_row + 1, max_r + 1):
            pn = ws.cell(row=r, column=pn_col).value if pn_col else None
            desc = ws.cell(row=r, column=desc_col).value if desc_col else ""
            tot = ws.cell(row=r, column=total_col).value if total_col else None
            spares = ws.cell(row=r, column=spares_col).value if spares_col else None
            if not pn:
                continue
            text = f"{pn} {desc}".upper()
            if any(k in text for k in KEYWORDS):
                print(f"    row {r:>4}  PN={str(pn).strip()!r:30}  total={tot!r:>6}  spares={spares!r:>4}  desc={desc!r}")

if __name__ == "__main__":
    main()
