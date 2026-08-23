from openpyxl import load_workbook
from pathlib import Path
P = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/Sales_BPA_NGT_Refresh_TO4_SalesBoM_Site_Packing_2026-05-23_00-33.xlsx")
wb = load_workbook(P, data_only=True)
ws = wb["BOM"]
print(f"max_col={ws.max_column}")
for c in range(1, ws.max_column+1):
    v7 = ws.cell(7, c).value
    v8 = ws.cell(8, c).value
    if v7 or v8:
        print(f"  col {c:>3}: r7={v7!r:35}  r8={v8!r}")
