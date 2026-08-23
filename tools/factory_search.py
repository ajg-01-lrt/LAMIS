"""Search Factory Live Inventory for NMA-8509 and the other shortfall PNs."""
from __future__ import annotations
import re
from pathlib import Path
from openpyxl import load_workbook

LIVE = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/Master_BPA_TO4_Factoy_Inventory.BOM.xlsx")

PNS = ["NMA-8509", "NMA8509", "ONS-SI-155", "1AB215120077", "3HE13998AA", "3HE11286AA"]

def norm(s): return re.sub(r"\s+", "", str(s or "")).upper()

wb = load_workbook(LIVE, data_only=True)
for needle in PNS:
    print(f"\n========== Searching for {needle!r} ==========")
    nn = norm(needle)
    hits = 0
    for ws in wb.worksheets:
        max_r = min(ws.max_row or 0, 5000)
        max_c = min(ws.max_column or 0, 200)
        for row in ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c, values_only=False):
            for cell in row:
                if cell.value is None: continue
                nv = norm(cell.value)
                if nn in nv:
                    rv = [(c, ws.cell(cell.row, c).value) for c in range(1, max_c+1)]
                    rv = [(c, v) for c, v in rv if v not in (None, "")]
                    print(f"  [{ws.title}] {cell.coordinate}: value={cell.value!r}")
                    print(f"     row cells: {rv[:20]}")
                    hits += 1
                    if hits >= 6:
                        break
            if hits >= 6: break
        if hits >= 6: break
    if hits == 0:
        print("  -- NO HITS --")
