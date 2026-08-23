from openpyxl import load_workbook
from pathlib import Path
SRC = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/ATLAS Build Material/TO4-Services-Prices-with-BoM-V9.1-03-05-26_BPA.xlsx")
wb = load_workbook(SRC, data_only=True)
ws = wb["TO4 BoM v8"]
# Dump cols 65-81 header rows 1-12 + data rows for 3HE03127AA (r33)
print("HEADERS rows 1-12, cols 65-81:")
for r in range(1, 13):
    cells = []
    for c in range(65, 82):
        v = ws.cell(r, c).value
        if v is not None:
            cells.append(f"col{c}={v!r}")
    if cells:
        print(f"  r{r}: " + " | ".join(cells)[:300])

print("\n--- 3HE03127AA row (r33) cols 65-81 ---")
for c in range(65, 82):
    v = ws.cell(33, c).value
    print(f"  col {c}: {v!r}")

print("\n--- A few other PN rows cols 65-81 ---")
for r in (10, 12, 15, 20, 25, 33, 40, 50, 60):
    pn = ws.cell(r, 2).value
    if not pn:
        continue
    print(f"\n  r{r} PN={pn!r}:")
    for c in range(65, 82):
        v = ws.cell(r, c).value
        if v not in (None, "", 0):
            print(f"    col {c}: {v!r}")
