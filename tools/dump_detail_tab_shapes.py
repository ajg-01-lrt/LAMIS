"""Dump structure of per-site detail tabs (one Sales, one Factory) so we
know the header row and column layout to parse them generically."""
from openpyxl import load_workbook
from pathlib import Path

SALES = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/Sales_BPA_NGT_Refresh_TO4_SalesBoM_Site_Packing_2026-05-23_00-33.xlsx")
LIVE  = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/Master_BPA_TO4_Factoy_Inventory.BOM.xlsx")

def dump(p, sheets, label, max_rows=22):
    wb = load_workbook(p, data_only=True)
    print(f"\n{'='*60}\n{label}: {p.name}\n{'='*60}")
    print(f"All sheet titles: {wb.sheetnames[:8]} ... ({len(wb.sheetnames)} total)")
    for sheet_name in sheets:
        if sheet_name not in wb.sheetnames:
            print(f"\n  -- {sheet_name!r} not present --")
            continue
        ws = wb[sheet_name]
        print(f"\n-- {sheet_name!r} (max_row={ws.max_row} max_col={ws.max_column}) --")
        for r in range(1, min(max_rows, (ws.max_row or 0)+1)+1):
            cells = []
            for c in range(1, min(ws.max_column or 0, 12)+1):
                v = ws.cell(r,c).value
                if v is not None:
                    cells.append(f"[{c}]{v!r}")
            if cells:
                print(f"  r{r:>2}: " + "  ".join(cells))

dump(SALES, ["ALBA", "Spares"], "SALES")
dump(LIVE,  ["MALN", "SPARES"], "FACTORY")
