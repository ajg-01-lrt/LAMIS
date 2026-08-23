from openpyxl import load_workbook
from pathlib import Path
SRC = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/ATLAS Build Material/TO4-Services-Prices-with-BoM-V9.1-03-05-26_BPA.xlsx")
wb = load_workbook(SRC, data_only=True)
ws = wb["TO4 BoM v8"]
print(f"TO4 BoM v8: {ws.max_row}x{ws.max_column}")
# Headers usually rows 1-10
for r in range(1, 12):
    cells = []
    for c in range(1, min(ws.max_column or 0, 90)+1):
        v = ws.cell(r, c).value
        if v is not None:
            cells.append(f"[{c}]{v!r}")
    if cells:
        print(f"  r{r}: " + "  ".join(cells)[:380])
