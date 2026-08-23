from openpyxl import load_workbook
from pathlib import Path
P = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/Sales_BPA_NGT_Refresh_TO4_SalesBoM_Site_Packing_2026-05-23_00-33.xlsx")
wb = load_workbook(P, data_only=True)
ws = wb["Spares"]
print(f"Spares: {ws.max_row}x{ws.max_column}")
for r in range(14, ws.max_row+1):
    cells = [str(ws.cell(r, c).value or "") for c in range(1, 5)]
    if any(cells):
        print(f"  r{r:>3}: B={cells[1]:<25} C={cells[2]:<55} D={cells[3]}")
