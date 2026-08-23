from openpyxl import load_workbook
from pathlib import Path
SRC = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/ATLAS Build Material/TO4-Services-Prices-with-BoM-V9.1-03-05-26_BPA.xlsx")
wb = load_workbook(SRC, data_only=True)
ws = wb["CBR_bpa version"]
print(f"CBR_bpa version: {ws.max_row}x{ws.max_column}")
# Headers around rows 1-15
for r in range(1, 16):
    cells = []
    for c in range(1, min(ws.max_column or 0, 40)+1):
        v = ws.cell(r, c).value
        if v is not None:
            cells.append(f"[{c}]{v!r}")
    if cells:
        print(f"  r{r}: " + "  ".join(cells)[:250])
print("\n--- Looking for 3HE03127AA ---")
for r in range(1, ws.max_row+1):
    for c in range(1, min(ws.max_column or 0, 40)+1):
        v = ws.cell(r, c).value
        if v and "3HE03127" in str(v):
            row_data = [(cc, ws.cell(r, cc).value) for cc in range(1, ws.max_column+1)]
            row_data = [(cc, v) for cc, v in row_data if v not in (None, "")]
            print(f"  r{r} col {c} = {v!r}")
            print(f"    full row: {row_data[:25]}")
            break
