"""Compare the Live Inventory tab's Total column values to what we'd get
by counting rows in per-site detail tabs (the ground truth).

If the recompute is double-counting (e.g. summing both site cells AND
rollup cells), we'll see Total > detail-tab count."""
from openpyxl import load_workbook
from pathlib import Path
from collections import defaultdict

COMPARE = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/Master_BPA_TO4_Factoy_Inventory.BOM.COMPARE.xlsx")
LIVE_SRC = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/Master_BPA_TO4_Factoy_Inventory.BOM.xlsx")

wb_cmp = load_workbook(COMPARE, data_only=True)
ws_inv = wb_cmp["Live Inventory"]
print(f"Live Inventory tab dims: {ws_inv.max_row}x{ws_inv.max_column}")

# Find the rightmost Total column on row 7
total_col = None
for c in range(ws_inv.max_column, 0, -1):
    v = ws_inv.cell(7, c).value
    if v and str(v).strip().lower() == "total":
        total_col = c
        break
print(f"Total col on Live Inventory: {total_col}")

# Read each PN row's Total cell vs sum of per-site cells (cols 4..total_col-1)
print("\nPN | Total cell | Sum of site cells | Ratio")
for r in range(10, min(ws_inv.max_row, 80) + 1):
    pn = ws_inv.cell(r, 2).value
    if not pn: continue
    total = ws_inv.cell(r, total_col).value
    sum_cells = 0
    for c in range(4, total_col):
        v = ws_inv.cell(r, c).value
        if isinstance(v, (int, float)):
            sum_cells += int(v)
    print(f"  {pn!r:25} total={total!r:>8}  sum_4_to_pre_total={sum_cells:>5}")

# Now ground-truth: count physical units in per-site detail tabs of the SOURCE
print("\n--- Ground-truth counts from Live Inventory SOURCE per-site tabs ---")
wb_src = load_workbook(LIVE_SRC, data_only=True)
TARGETS = ("3HE04824AA", "3HE04823AA", "3HE12546AA", "3HE07943AA", "3HE11278AA")
counts = defaultdict(int)
for ws in wb_src.worksheets:
    if ws.title.upper() in ("BOM", "SUMMARY"): continue
    max_r = min(ws.max_row or 0, 5000)
    for r in range(15, max_r+1):
        pn = ws.cell(r, 4).value  # Factory layout: PN at col D
        if pn and str(pn).strip() in TARGETS:
            counts[str(pn).strip()] += 1
for pn, n in sorted(counts.items()):
    print(f"  {pn}: {n} physical units (detail tabs)")
