"""Dump the entire Shortfall tab so we can see what rows look suspicious."""
from openpyxl import load_workbook
from pathlib import Path
P = Path(r"C:/Users/ZackerySimino/OneDrive - LightRiver Technologies Inc/Desktop/Inventory and Packing slips/Testing/Master_BPA_TO4_Factoy_Inventory.BOM.COMPARE.xlsx")
wb = load_workbook(P, data_only=True)
print("Sheets:", wb.sheetnames)
ws = wb["Shortfall"]
print(f"\nShortfall: {ws.max_row}x{ws.max_column}")
for r in range(1, ws.max_row + 1):
    cells = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(r, c).value
        cells.append("" if v is None else str(v))
    if any(cells):
        print(f"  r{r:>3}: " + " | ".join(f"{x:<35}" if c==2 else f"{x:<10}" for c,x in enumerate(cells, start=1)))
